from __future__ import annotations

import hashlib
import json
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

import models


MONTHS_PT = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}


@dataclass
class ParsedLancamento:
    linha_planilha: int
    secao: str | None
    indice_item: int | None
    descricao: str | None
    valor: float | None
    observacao: str | None
    raw_row: dict[str, Any]


@dataclass
class ParsedExtrato:
    planilha_nome: str
    dealer_nome: str | None
    dealer_account: str | None
    total_expenditure: float | None
    total_income: float | None
    credit_balance: float | None
    total_prepayment: float | None
    raw_rows: list[dict[str, Any]]
    lancamentos: list[ParsedLancamento]


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text


def _as_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text.replace(".", "").replace(",", "."))
    except ValueError:
        return None


def _to_int(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _extract_competencia(file_path: Path) -> tuple[str | None, int | None, int | None, date | None]:
    stem = file_path.stem.strip()
    parts = stem.split()
    if len(parts) < 2:
        return None, None, None, None

    month_name_raw = parts[0]
    year_raw = parts[-1]
    month_norm = _normalize_text(month_name_raw)
    month = MONTHS_PT.get(month_norm)
    try:
        year = int(year_raw)
    except ValueError:
        year = None

    competencia = f"{month_name_raw} {year_raw}" if month_name_raw and year_raw else None
    competencia_data = date(year, month, 1) if (month and year) else None
    return competencia, month, year, competencia_data


def parse_credito_file(file_path: Path) -> ParsedExtrato:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    dealer_nome = None
    dealer_account = None
    total_expenditure = None
    total_income = None
    credit_balance = None
    total_prepayment = None
    current_section = None

    raw_rows: list[dict[str, Any]] = []
    lancamentos: list[ParsedLancamento] = []

    for row_idx in range(1, ws.max_row + 1):
        row_values = [ws.cell(row_idx, col).value for col in range(1, 9)]
        if not any(v is not None and str(v).strip() != "" for v in row_values):
            continue

        row_obj = {
            "linha": row_idx,
            "A": row_values[0],
            "B": row_values[1],
            "C": row_values[2],
            "D": row_values[3],
            "E": row_values[4],
            "F": row_values[5],
            "G": row_values[6],
            "H": row_values[7],
        }
        raw_rows.append(row_obj)

        label_b = _as_text(row_values[1])
        norm_b = _normalize_text(label_b)
        idx_item = _to_int(row_values[2])
        descricao = _as_text(row_values[3])
        valor = _to_float(row_values[4])
        observacao = _as_text(row_values[5])

        if norm_b == "dealer":
            dealer_nome = _as_text(row_values[2])

        label_e = _as_text(row_values[4])
        if "dealer account" in _normalize_text(label_e):
            dealer_account = _as_text(row_values[5])

        if norm_b.startswith("expenditure item"):
            current_section = "expenditure"
        elif norm_b.startswith("income item"):
            current_section = "income"
        elif norm_b.startswith("pre-payment"):
            current_section = "pre_payment"

        if norm_b.startswith("total expenditure"):
            total_expenditure = valor
            lancamentos.append(
                ParsedLancamento(
                    linha_planilha=row_idx,
                    secao="summary",
                    indice_item=None,
                    descricao=label_b,
                    valor=valor,
                    observacao=observacao,
                    raw_row=row_obj,
                )
            )
            continue

        if norm_b.startswith("total income"):
            total_income = valor
            lancamentos.append(
                ParsedLancamento(
                    linha_planilha=row_idx,
                    secao="summary",
                    indice_item=None,
                    descricao=label_b,
                    valor=valor,
                    observacao=observacao,
                    raw_row=row_obj,
                )
            )
            continue

        if norm_b.startswith("credit balance"):
            credit_balance = valor
            lancamentos.append(
                ParsedLancamento(
                    linha_planilha=row_idx,
                    secao="summary",
                    indice_item=None,
                    descricao=label_b,
                    valor=valor,
                    observacao=observacao,
                    raw_row=row_obj,
                )
            )
            continue

        if norm_b.startswith("total pre-payment"):
            total_prepayment = valor
            lancamentos.append(
                ParsedLancamento(
                    linha_planilha=row_idx,
                    secao="summary",
                    indice_item=None,
                    descricao=label_b,
                    valor=valor,
                    observacao=observacao,
                    raw_row=row_obj,
                )
            )
            continue

        if idx_item is not None and descricao is not None:
            lancamentos.append(
                ParsedLancamento(
                    linha_planilha=row_idx,
                    secao=current_section,
                    indice_item=idx_item,
                    descricao=descricao,
                    valor=valor,
                    observacao=observacao,
                    raw_row=row_obj,
                )
            )

    wb.close()
    return ParsedExtrato(
        planilha_nome=ws.title,
        dealer_nome=dealer_nome,
        dealer_account=dealer_account,
        total_expenditure=total_expenditure,
        total_income=total_income,
        credit_balance=credit_balance,
        total_prepayment=total_prepayment,
        raw_rows=raw_rows,
        lancamentos=lancamentos,
    )


def sync_credito_files_to_db(db: Session, credito_dir: Path) -> dict[str, Any]:
    credito_dir = credito_dir.resolve()
    credito_dir.mkdir(parents=True, exist_ok=True)

    summary = {
        "processed": 0,
        "imported": 0,
        "already_present": 0,
        "failed": 0,
        "errors": [],
    }

    files = sorted([p for p in credito_dir.glob("*.xlsx") if p.is_file()], key=lambda p: p.name.lower())
    for file_path in files:
        summary["processed"] += 1
        rel = file_path.relative_to(credito_dir).as_posix()
        sha = _sha256_file(file_path)
        size = file_path.stat().st_size

        existing = (
            db.query(models.CreditoExtrato)
            .filter(
                models.CreditoExtrato.arquivo_caminho_relativo == rel,
                models.CreditoExtrato.arquivo_sha256 == sha,
            )
            .first()
        )
        if existing:
            summary["already_present"] += 1
            continue

        try:
            parsed = parse_credito_file(file_path)
            competencia_nome, competencia_mes, competencia_ano, competencia_data = _extract_competencia(file_path)

            extrato = models.CreditoExtrato(
                arquivo_nome=file_path.name,
                arquivo_caminho_relativo=rel,
                arquivo_sha256=sha,
                arquivo_tamanho_bytes=size,
                competencia_nome=competencia_nome,
                competencia_mes=competencia_mes,
                competencia_ano=competencia_ano,
                competencia_data=competencia_data,
                planilha_nome=parsed.planilha_nome,
                dealer_nome=parsed.dealer_nome,
                dealer_account=parsed.dealer_account,
                total_expenditure=parsed.total_expenditure,
                total_income=parsed.total_income,
                credit_balance=parsed.credit_balance,
                total_prepayment=parsed.total_prepayment,
                raw_rows_json=json.dumps(parsed.raw_rows, ensure_ascii=False),
            )
            db.add(extrato)
            db.flush()

            for lanc in parsed.lancamentos:
                db.add(
                    models.CreditoLancamento(
                        extrato_id=extrato.id,
                        linha_planilha=lanc.linha_planilha,
                        secao=lanc.secao,
                        indice_item=lanc.indice_item,
                        descricao=lanc.descricao,
                        valor=lanc.valor,
                        observacao=lanc.observacao,
                        raw_row_json=json.dumps(lanc.raw_row, ensure_ascii=False),
                    )
                )

            db.commit()
            summary["imported"] += 1
        except Exception as exc:
            db.rollback()
            summary["failed"] += 1
            summary["errors"].append({"arquivo": file_path.name, "erro": str(exc)})

    return summary

